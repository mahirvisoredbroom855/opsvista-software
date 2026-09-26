"""
Tests for app/features/rag_chatbot/ingestion_common.py — the shared
extraction/chunking logic used by both build_local_index.py and
build_drive_index.py.
"""
from __future__ import annotations

import io

from app.features.rag_chatbot.ingestion_common import (
    chunk_text,
    department_for_folder,
    extract_table_chunks_from_xlsx_bytes,
    extract_text_from_docx_bytes,
    extract_text_from_file,
    extract_text_from_pdf_bytes,
    make_chunk_documents,
    make_table_chunk_documents,
)


def test_department_for_folder_matches_partial_names():
    assert department_for_folder("Md. Mozammel Haque (HR)") == "HR"
    assert department_for_folder("Md. Mizanur Rahman (PTIL)") == "Finance"
    assert department_for_folder("Some Unrelated Folder") == "Unknown"


def test_chunk_text_respects_max_chars():
    text = "\n\n".join(f"Paragraph {i} " + ("x" * 100) for i in range(10))
    chunks = chunk_text(text, max_chars=300)
    assert len(chunks) > 1
    # Every chunk should be reasonably close to the budget, not wildly over.
    assert all(len(c) <= 300 + 120 for c in chunks)  # +1 paragraph slack, same logic as chunk_text


def test_chunk_text_handles_short_text():
    chunks = chunk_text("Just one short paragraph.")
    assert chunks == ["Just one short paragraph."]


def test_extract_text_from_docx_bytes_reads_paragraphs_and_tables():
    from docx import Document

    doc = Document()
    doc.add_heading("Test Policy", level=1)
    doc.add_paragraph("This is a test paragraph.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Field"
    table.rows[0].cells[1].text = "Value"

    buf = io.BytesIO()
    doc.save(buf)

    text = extract_text_from_docx_bytes(buf.getvalue())
    assert "Test Policy" in text
    assert "This is a test paragraph." in text
    assert "Field | Value" in text


def test_extract_text_from_pdf_bytes_reads_text():
    from fpdf import FPDF

    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.cell(0, 10, "Hello from a test PDF")
    raw = bytes(pdf.output())

    text = extract_text_from_pdf_bytes(raw)
    assert "Hello from a test PDF" in text


def test_extract_table_chunks_from_xlsx_bytes_preserves_structure():
    import pandas as pd

    df = pd.DataFrame(
        {
            "Line Item": ["Revenue", "COGS", "Profit"],
            "Q1": [100, 40, 60],
            "Q2": [120, 45, 75],
        }
    )
    buf = io.BytesIO()
    df.to_excel(buf, index=False, sheet_name="P&L")
    raw = buf.getvalue()

    chunks = extract_table_chunks_from_xlsx_bytes(raw, rows_per_group=10)
    assert len(chunks) == 1
    table = chunks[0]["table"]
    assert table["sheet"] == "P&L"
    assert table["columns"] == ["Line Item", "Q1", "Q2"]
    assert table["rows"][0][0] == "Revenue"
    assert "Revenue" in chunks[0]["text"]


def test_extract_text_from_file_dispatches_by_extension():
    assert extract_text_from_file(b"hello", "notes.txt") == "hello"
    assert extract_text_from_file(b"# Title", "notes.md") == "# Title"
    assert extract_text_from_file(b"whatever", "unsupported.xyz") == ""


def test_make_chunk_documents_shapes_metadata_correctly():
    docs = make_chunk_documents(
        "Paragraph one.\n\nParagraph two.",
        id_prefix="test_doc",
        file_name="policy.md",
        title="Policy",
        department="HR",
        owner_folder="Some Folder",
        source="local_seed",
    )
    assert len(docs) >= 1
    first = docs[0]
    assert first["id"] == "test_doc_0"
    assert first["metadata"]["department"] == "HR"
    assert first["metadata"]["chunk_index"] == 0
    assert first["text"].startswith("Policy\n\n")


def test_make_table_chunk_documents_carries_table_metadata():
    table_chunks = [
        {"text": "Sheet 'X': ...", "table": {"sheet": "X", "columns": ["A"], "rows": [["1"]]}},
    ]
    docs = make_table_chunk_documents(
        table_chunks,
        id_prefix="test_doc",
        file_name="data.xlsx",
        title="Data",
        department="Finance",
        owner_folder="Some Folder",
        source="google_drive",
    )
    assert docs[0]["metadata"]["table"]["sheet"] == "X"
    assert docs[0]["metadata"]["source"] == "google_drive"
