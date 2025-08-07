# Excel Processing Engine for Finance Files
import openpyxl
from openpyxl import Workbook
from io import BytesIO
from typing import List, Dict, Optional, Tuple, Any
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
import logging
import re

from .models import RawDriveFinance, FactFinance, FinanceFileProcessing
from sqlalchemy.orm import Session

logger = logging.getLogger(__name__)

class ExcelProcessor:
    """Smart Excel file processor for finance data"""
    
    def __init__(self, db_session: Session):
        self.db = db_session
        self.processing_stats = {
            'rows_processed': 0,
            'rows_successful': 0,
            'rows_failed': 0,
            'errors': []
        }
    
    # =====================================
    # FILE TYPE DETECTION FUNCTIONS
    # =====================================
    
    def detect_file_type(self, file_name: str, file_content: bytes) -> str:
        """Detect Excel file type based on name and structure"""
        logger.info(f"🔍 Detecting file type for: {file_name}")
        
        # Method 1: Name-based detection (most reliable)
        if "Cash Book" in file_name:
            logger.info("📊 Detected: Cash Book file")
            return "cash_book"
        elif "Party Due Bill" in file_name or "Due Bill" in file_name:
            logger.info("💰 Detected: Party Due Bill file")
            return "party_due_bill" 
        elif "Expenditure" in file_name or "PTIL Expenditure" in file_name:
            logger.info("📈 Detected: Expenditure Summary file")
            return "expenditure_summary"
        
        # Method 2: Content-based detection (fallback)
        try:
            workbook = openpyxl.load_workbook(BytesIO(file_content))
            sheet = workbook.active
            
            # Look for specific column patterns in first few rows
            for row_num in range(1, 6):
                headers = [str(cell.value or "") for cell in sheet[row_num]]
                headers_text = " ".join(headers).lower()
                
                if "cash received" in headers_text:
                    logger.info("📊 Detected: Cash Book (by content)")
                    return "cash_book"
                elif "party name" in headers_text and "due bill" in headers_text:
                    logger.info("💰 Detected: Party Due Bill (by content)")
                    return "party_due_bill"
                elif "expenditure" in headers_text or ("year" in headers_text and "profit" in headers_text):
                    logger.info("📈 Detected: Expenditure Summary (by content)")
                    return "expenditure_summary"
            
        except Exception as e:
            logger.warning(f"⚠️ Content detection failed: {str(e)}")
        
        logger.warning("❓ Unknown file type")
        return "unknown"
    
    def find_header_row(self, worksheet) -> int:
        """Find the row containing column headers"""
        logger.info("🔍 Looking for header row...")
        
        # Common header indicators for finance files
        header_indicators = [
            "date", "cash", "amount", "party", "bill", 
            "total", "expenditure", "year", "description",
            "particulars", "sl", "qty", "advance"
        ]
        
        # Check first 10 rows for headers
        for row_num in range(1, 11):
            row_values = [str(cell.value or "").lower() for cell in worksheet[row_num]]
            
            # Count how many header indicators we find
            matches = sum(1 for indicator in header_indicators 
                         if any(indicator in val for val in row_values))
            
            # If we find 2+ indicators, this is likely the header row
            if matches >= 2:
                logger.info(f"✅ Found header row: {row_num} (matched {matches} indicators)")
                return row_num
        
        logger.info("📍 Using default header row: 2")
        return 2  # Default to row 2 if no clear headers found
    
    # =====================================
    # DATE PARSING FUNCTIONS
    # =====================================
    
    def smart_date_parsing(self, date_value) -> Optional[date]:
        """Parse dates in multiple formats from Excel"""
        if date_value is None or date_value == "":
            return None
        
        # Handle Excel date numbers (e.g., 44927.0 = 2023-01-01)
        if isinstance(date_value, (int, float)):
            try:
                # Excel epoch starts 1900-01-01 (with leap year bug)
                excel_epoch = datetime(1899, 12, 30)  # Excel's actual epoch
                parsed_date = (excel_epoch + timedelta(days=int(date_value))).date()
                
                # Sanity check: date should be reasonable
                if date(1900, 1, 1) <= parsed_date <= date(2100, 12, 31):
                    return parsed_date
            except Exception:
                pass
        
        # Handle string dates in various formats
        date_formats = [
            "%d.%m.%Y",     # 01.06.2025
            "%d/%m/%Y",     # 01/06/2025  
            "%Y-%m-%d",     # 2025-06-01
            "%m/%d/%Y",     # 06/01/2025
            "%d-%m-%Y",     # 01-06-2025
            "%Y.%m.%d",     # 2025.06.01
        ]
        
        date_str = str(date_value).strip()
        
        for fmt in date_formats:
            try:
                parsed_date = datetime.strptime(date_str, fmt).date()
                
                # Sanity check
                if date(1900, 1, 1) <= parsed_date <= date(2100, 12, 31):
                    return parsed_date
            except ValueError:
                continue
        
        # Last resort: try to extract numbers and guess format
        numbers = re.findall(r'\d+', date_str)
        
        if len(numbers) == 3:
            # Try different interpretations of the 3 numbers
            interpretations = [
                (int(numbers[0]), int(numbers[1]), int(numbers[2])),  # day, month, year
                (int(numbers[2]), int(numbers[1]), int(numbers[0])),  # year, month, day
            ]
            
            for day, month, year in interpretations:
                try:
                    # Handle 2-digit years
                    if year < 100:
                        year += 2000 if year < 50 else 1900
                    
                    if 1 <= month <= 12 and 1 <= day <= 31 and 1900 <= year <= 2100:
                        return date(year, month, day)
                except ValueError:
                    continue
        
        return None  # Could not parse
    
    # =====================================
    # NUMBER PARSING FUNCTIONS
    # =====================================
    
    def safe_decimal(self, value) -> Optional[Decimal]:
        """Safely parse numeric values to Decimal"""
        if value is None or value == "":
            return None
        
        try:
            # Handle different numeric formats
            if isinstance(value, (int, float)):
                return Decimal(str(value))
            
            # Clean string values
            clean_value = str(value).strip()
            
            # Remove common non-numeric characters
            clean_value = re.sub(r'[,\s]', '', clean_value)  # Remove commas and spaces
            clean_value = re.sub(r'[^\d.-]', '', clean_value)  # Keep only digits, dots, minus
            
            if clean_value == "" or clean_value == "-":
                return None
            
            return Decimal(clean_value)
            
        except (InvalidOperation, ValueError):
            return None
    
    def safe_int(self, value) -> Optional[int]:
        """Safely parse integer values"""
        if value is None or value == "":
            return None
        
        try:
            if isinstance(value, int):
                return value
            elif isinstance(value, float):
                return int(value)
            else:
                clean_value = str(value).strip()
                clean_value = re.sub(r'[^\d-]', '', clean_value)
                
                if clean_value == "" or clean_value == "-":
                    return None
                
                return int(clean_value)
        except (ValueError, TypeError):
            return None
    
    # =====================================
    # MAIN PROCESSING FUNCTIONS
    # =====================================
    
    def process_file(self, file_content: bytes, file_name: str, processing_record_id: int) -> Dict[str, Any]:
        """Main entry point for processing Excel files"""
        try:
            logger.info(f"🔄 Starting processing: {file_name}")
            
            # Reset processing stats
            self.processing_stats = {
                'rows_processed': 0,
                'rows_successful': 0,
                'rows_failed': 0,
                'errors': []
            }
            
            # Detect file type
            file_type = self.detect_file_type(file_name, file_content)
            
            if file_type == "unknown":
                raise ValueError(f"Unknown file type: {file_name}")
            
            # Load Excel workbook
            workbook = openpyxl.load_workbook(BytesIO(file_content))
            
            # Process each sheet
            for sheet_name in workbook.sheetnames:
                logger.info(f"📄 Processing sheet: {sheet_name}")
                worksheet = workbook[sheet_name]
                
                # Process based on file type
                if file_type == "cash_book":
                    self.process_cash_book_sheet(worksheet, file_name, sheet_name, processing_record_id)
                elif file_type == "party_due_bill":
                    self.process_due_bill_sheet(worksheet, file_name, sheet_name, processing_record_id)
                elif file_type == "expenditure_summary":
                    self.process_expenditure_sheet(worksheet, file_name, sheet_name, processing_record_id)
            
            logger.info(f"✅ Processing completed: {file_name}")
            logger.info(f"   Processed: {self.processing_stats['rows_processed']} rows")
            logger.info(f"   Successful: {self.processing_stats['rows_successful']} rows")
            logger.info(f"   Failed: {self.processing_stats['rows_failed']} rows")
            
            return {
                'status': 'success',
                'file_type': file_type,
                'stats': self.processing_stats
            }
            
        except Exception as e:
            logger.error(f"❌ Processing failed for {file_name}: {str(e)}")
            return {
                'status': 'error',
                'error': str(e),
                'stats': self.processing_stats
            }
    
    # =====================================
    # CASH BOOK PROCESSING
    # =====================================
    
    def process_cash_book_sheet(self, worksheet, file_name: str, sheet_name: str, processing_record_id: int):
        """Process Cash Book Excel sheet"""
        logger.info(f"💰 Processing Cash Book sheet: {sheet_name}")
        
        # Find header row
        header_row = self.find_header_row(worksheet)
        headers = [str(cell.value or "") for cell in worksheet[header_row]]
        
        logger.info(f"📋 Headers found: {headers}")
        
        # Map column indices
        column_mapping = self.map_cash_book_columns(headers)
        logger.info(f"🗺️ Column mapping: {column_mapping}")
        
        # Process data rows
        for row_num in range(header_row + 1, worksheet.max_row + 1):
            self.processing_stats['rows_processed'] += 1
            
            try:
                row = worksheet[row_num]
                
                # Extract data using column mapping
                transaction_date = self.smart_date_parsing(
                    row[column_mapping.get('date', 0)].value if column_mapping.get('date') is not None else None
                )
                
                if not transaction_date:
                    continue  # Skip rows without valid dates
                
                # Extract amounts from different columns
                cash_received = self.safe_decimal(
                    row[column_mapping.get('cash_received', 1)].value if column_mapping.get('cash_received') is not None else None
                )
                factory_amount = self.safe_decimal(
                    row[column_mapping.get('factory', 2)].value if column_mapping.get('factory') is not None else None
                )
                particulars = str(row[column_mapping.get('particulars', 6)].value or "")
                
                # Save raw data first
                raw_record = RawDriveFinance(
                    ingestion_time=datetime.now(),
                    source_file=file_name,
                    sheet_name=sheet_name,
                    row_number=row_num,
                    transaction_date=transaction_date,
                    cash_received=cash_received,
                    factory_amount=factory_amount,
                    particulars=particulars,
                    raw_data={
                        'headers': headers,
                        'row_data': [str(cell.value) for cell in row],
                        'processing_id': processing_record_id
                    },
                    processing_status='pending'
                )
                
                self.db.add(raw_record)
                self.db.flush()  # Get the ID
                
                # Transform to fact table if we have meaningful data
                if cash_received and cash_received > 0:
                    fact_record = FactFinance(
                        transaction_date=transaction_date,
                        transaction_type='cash_receipt',
                        amount_bdt=float(cash_received),
                        amount_original=float(cash_received),
                        currency_original='BDT',
                        description=particulars,
                        transaction_category='daily_cash_flow',
                        source_reference=f"{file_name}#{sheet_name}#{row_num}",
                        data_quality_score=self.calculate_quality_score(raw_record),
                        # raw_data_id removed - not needed
                    )
                    self.db.add(fact_record)
                
                # Mark raw record as processed
                raw_record.processing_status = 'completed'
                self.processing_stats['rows_successful'] += 1
                
            except Exception as e:
                self.processing_stats['rows_failed'] += 1
                self.processing_stats['errors'].append(f"Row {row_num}: {str(e)}")
                logger.warning(f"⚠️ Failed to process row {row_num}: {str(e)}")
        
        self.db.commit()
    
    def map_cash_book_columns(self, headers: List[str]) -> Dict[str, Optional[int]]:
        """Map Cash Book column headers to indices"""
        mapping = {}
        
        for i, header in enumerate(headers):
            header_lower = header.lower().strip()
            
            if 'date' in header_lower:
                mapping['date'] = i
            elif 'cash received' in header_lower or 'cash' in header_lower:
                mapping['cash_received'] = i
            elif 'factory' in header_lower and 'cr' not in header_lower:
                mapping['factory'] = i
            elif 'particulars' in header_lower or 'description' in header_lower:
                mapping['particulars'] = i
            elif 'total' in header_lower:
                mapping['total'] = i
        
        return mapping
    
    # =====================================
    # QUALITY SCORING
    # =====================================
    
    def calculate_quality_score(self, raw_record) -> float:
        """Calculate data quality score (0.0 to 1.0)"""
        score = 1.0
        
        # Deduct points for missing data
        if not raw_record.transaction_date:
            score -= 0.3
        if not raw_record.cash_received and not raw_record.factory_amount:
            score -= 0.4
        if not raw_record.particulars or raw_record.particulars.strip() == "":
            score -= 0.2
        
        # Bonus points for rich data
        if raw_record.particulars and len(raw_record.particulars.strip()) > 10:
            score += 0.1
        
        return max(0.0, min(1.0, score))

    # =====================================
    # COLUMN MAPPING METHODS (Missing ones)
    # =====================================
    
    def map_due_bill_columns(self, headers: List[str]) -> Dict[str, Optional[int]]:
        """Map Due Bill column headers to indices"""
        mapping = {}
        
        for i, header in enumerate(headers):
            header_lower = header.lower().strip()
            
            if 'sl' in header_lower and len(header_lower) <= 3:
                mapping['sl'] = i
            elif 'date' in header_lower:
                mapping['date'] = i
            elif 'party' in header_lower and 'name' in header_lower:
                mapping['party_name'] = i
            elif 'party' in header_lower and 'name' not in header_lower:
                mapping['party_name'] = i
            elif 'description' in header_lower or 'particulars' in header_lower:
                mapping['description'] = i
            elif 'bill' in header_lower and 'no' in header_lower:
                mapping['bill_no'] = i
            elif 'qty' in header_lower or 'quantity' in header_lower:
                mapping['quantity'] = i
            elif 'unit' in header_lower and 'price' in header_lower:
                mapping['unit_price'] = i
            elif 'total' in header_lower and 'amount' in header_lower:
                mapping['total_amount'] = i
            elif 'advance' in header_lower:
                mapping['advance'] = i
            elif 'due' in header_lower and 'bill' in header_lower:
                mapping['due_bill'] = i
            elif 'amount' in header_lower and 'total' not in header_lower and 'due' not in header_lower:
                mapping['total_amount'] = i
        
        return mapping
    
    def map_expenditure_columns(self, headers: List[str]) -> Dict[str, Optional[int]]:
        """Map Expenditure column headers to indices"""
        mapping = {}
        
        for i, header in enumerate(headers):
            header_lower = header.lower().strip()
            
            if 'year' in header_lower:
                mapping['year'] = i
            elif 'expenditure' in header_lower or 'expense' in header_lower:
                mapping['expenditure'] = i
            elif 'revenue' in header_lower or 'income' in header_lower or 'earning' in header_lower:
                mapping['revenue'] = i
            elif 'profit' in header_lower or 'loss' in header_lower:
                mapping['profit_loss'] = i
            elif 'total' in header_lower and ('expenditure' in header_lower or 'expense' in header_lower):
                mapping['expenditure'] = i
            elif 'net' in header_lower and ('profit' in header_lower or 'income' in header_lower):
                mapping['profit_loss'] = i
        
        return mapping

    # =====================================
    # ENSURE PROCESSING METHODS ARE COMPLETE
    # =====================================
    
    def calculate_quality_score(self, raw_record) -> float:
        """Calculate data quality score (0.0 to 1.0) - Enhanced version"""
        score = 1.0
        
        # Deduct points for missing data
        if not hasattr(raw_record, 'transaction_date') or not raw_record.transaction_date:
            score -= 0.3
        
        # Check for amount data
        amount_fields = ['cash_received', 'total_amount', 'amount_bdt', 'factory_amount']
        has_amount = any(getattr(raw_record, field, None) for field in amount_fields)
        if not has_amount:
            score -= 0.4
        
        # Check for description data
        desc_fields = ['particulars', 'description', 'party_name']
        has_description = any(getattr(raw_record, field, None) and str(getattr(raw_record, field, '')).strip() for field in desc_fields)
        if not has_description:
            score -= 0.2
        
        # Bonus points for rich data
        if has_description:
            desc_text = str(getattr(raw_record, 'particulars', '') or getattr(raw_record, 'description', '') or '')
            if len(desc_text.strip()) > 10:
                score += 0.1
        
        return max(0.0, min(1.0, score))

    # =====================================
    # COMPLETE DUE BILL PROCESSING METHOD
    # =====================================
    
    def process_due_bill_sheet(self, worksheet, file_name: str, sheet_name: str, processing_record_id: int):
        """Process Party Due Bill Excel sheet"""
        logger.info(f"💰 Processing Due Bill sheet: {sheet_name}")
        
        # Find header row (Due Bill files often have headers around row 5)
        header_row = self.find_header_row(worksheet)
        headers = [str(cell.value or "") for cell in worksheet[header_row]]
        
        logger.info(f"📋 Headers found: {headers}")
        
        # Map column indices for due bill format
        column_mapping = self.map_due_bill_columns(headers)
        logger.info(f"🗺️ Column mapping: {column_mapping}")
        
        # Process data rows
        for row_num in range(header_row + 1, worksheet.max_row + 1):
            self.processing_stats['rows_processed'] += 1
            
            try:
                row = worksheet[row_num]
                
                # Extract data using column mapping
                transaction_date = self.smart_date_parsing(
                    row[column_mapping.get('date', 3)].value if column_mapping.get('date') is not None else None
                )
                
                party_name = str(row[column_mapping.get('party_name', 2)].value or "").strip()
                total_amount = self.safe_decimal(
                    row[column_mapping.get('total_amount', 7)].value if column_mapping.get('total_amount') is not None else None
                )
                due_bill = self.safe_decimal(
                    row[column_mapping.get('due_bill', 9)].value if column_mapping.get('due_bill') is not None else None
                )
                description = str(row[column_mapping.get('description', 1)].value or "").strip()
                advance_amount = self.safe_decimal(
                    row[column_mapping.get('advance', 8)].value if column_mapping.get('advance') is not None else None
                )
                
                # Skip rows without meaningful data
                if not party_name and not total_amount:
                    continue
                
                # Save raw data first
                raw_record = RawDriveFinance(
                    ingestion_time=datetime.now(),
                    source_file=file_name,
                    sheet_name=sheet_name,
                    row_number=row_num,
                    transaction_date=transaction_date,
                    party_name=party_name,
                    total_amount=total_amount,
                    due_bill=due_bill,
                    description=description,
                    advance_tk=advance_amount,
                    raw_data={
                        'headers': headers,
                        'row_data': [str(cell.value) for cell in row],
                        'processing_id': processing_record_id
                    },
                    processing_status='pending'
                )
                
                self.db.add(raw_record)
                self.db.flush()  # Get the ID
                
                # Transform to fact table if we have meaningful data
                if total_amount and total_amount > 0:
                    fact_record = FactFinance(
                        transaction_date=transaction_date or date.today(),
                        transaction_type='due_bill',
                        amount_bdt=float(total_amount),
                        amount_original=float(total_amount),
                        currency_original='BDT',
                        party_name=party_name,
                        description=description,
                        transaction_category='supplier_payment',
                        total_amount=float(total_amount) if total_amount else None,
                        amount_due=float(due_bill) if due_bill else None,
                        advance_paid=float(advance_amount) if advance_amount else None,
                        source_reference=f"{file_name}#{sheet_name}#{row_num}",
                        data_quality_score=self.calculate_quality_score(raw_record)
                    )
                    self.db.add(fact_record)
                
                # Mark raw record as processed
                raw_record.processing_status = 'completed'
                self.processing_stats['rows_successful'] += 1
                
            except Exception as e:
                self.processing_stats['rows_failed'] += 1
                self.processing_stats['errors'].append(f"Row {row_num}: {str(e)}")
                logger.warning(f"⚠️ Failed to process row {row_num}: {str(e)}")
        
        self.db.commit()
    
    # =====================================
    # COMPLETE EXPENDITURE PROCESSING METHOD
    # =====================================
    
    def process_expenditure_sheet(self, worksheet, file_name: str, sheet_name: str, processing_record_id: int):
        """Process Expenditure Summary Excel sheet"""
        logger.info(f"📈 Processing Expenditure sheet: {sheet_name}")
        
        # Find header row
        header_row = self.find_header_row(worksheet)
        headers = [str(cell.value or "") for cell in worksheet[header_row]]
        
        logger.info(f"📋 Headers found: {headers}")
        
        # Map column indices for expenditure format
        column_mapping = self.map_expenditure_columns(headers)
        logger.info(f"🗺️ Column mapping: {column_mapping}")
        
        # Process data rows
        for row_num in range(header_row + 1, worksheet.max_row + 1):
            self.processing_stats['rows_processed'] += 1
            
            try:
                row = worksheet[row_num]
                
                # Extract year and amounts
                year_value = self.safe_int(
                    row[column_mapping.get('year', 0)].value if column_mapping.get('year') is not None else None
                )
                
                expenditure = self.safe_decimal(
                    row[column_mapping.get('expenditure', 1)].value if column_mapping.get('expenditure') is not None else None
                )
                
                revenue = self.safe_decimal(
                    row[column_mapping.get('revenue', 2)].value if column_mapping.get('revenue') is not None else None
                )
                
                profit_loss = self.safe_decimal(
                    row[column_mapping.get('profit_loss', 3)].value if column_mapping.get('profit_loss') is not None else None
                )
                
                # Skip rows without year or if it's a total row
                if not year_value or year_value < 2000 or year_value > 2100:
                    continue
                
                # Create transaction date (use Dec 31 of the year)
                transaction_date = date(year_value, 12, 31)
                
                # Save raw data first
                raw_record = RawDriveFinance(
                    ingestion_time=datetime.now(),
                    source_file=file_name,
                    sheet_name=sheet_name,
                    row_number=row_num,
                    transaction_date=transaction_date,
                    year=year_value,
                    total_expenditure=expenditure,
                    total_revenue=revenue,
                    profit_loss=profit_loss,
                    raw_data={
                        'headers': headers,
                        'row_data': [str(cell.value) for cell in row],
                        'processing_id': processing_record_id
                    },
                    processing_status='pending'
                )
                
                self.db.add(raw_record)
                self.db.flush()  # Get the ID
                
                # Transform to fact table for expenditure
                if expenditure and expenditure > 0:
                    fact_record = FactFinance(
                        transaction_date=transaction_date,
                        transaction_type='yearly_summary',
                        amount_bdt=float(expenditure),
                        amount_original=float(expenditure),
                        currency_original='BDT',
                        description=f"Year {year_value} expenditure summary",
                        transaction_category='yearly_expenditure',
                        total_expenditure=float(expenditure) if expenditure else None,
                        total_revenue=float(revenue) if revenue else None,
                        profit_loss=float(profit_loss) if profit_loss else None,
                        source_reference=f"{file_name}#{sheet_name}#{row_num}",
                        data_quality_score=self.calculate_quality_score(raw_record)
                    )
                    self.db.add(fact_record)
                
                # Mark raw record as processed
                raw_record.processing_status = 'completed'
                self.processing_stats['rows_successful'] += 1
                
            except Exception as e:
                self.processing_stats['rows_failed'] += 1
                self.processing_stats['errors'].append(f"Row {row_num}: {str(e)}")
                logger.warning(f"⚠️ Failed to process row {row_num}: {str(e)}")
        
        self.db.commit()
