"""
This file has the shared steps for turning a raw file — a PDF, Word
doc, or Excel sheet — into small, searchable pieces of text. Both ways
of building the search index (from a local test folder, or from real
Google Drive) use these exact same steps, so a document is chunked and
labeled the same way no matter where it came from.

Shared helpers used by build_local_index.py and build_drive_index.py so both
ingestion paths (local folder / real Google Drive) chunk and classify
documents identically. Deliberately simple — this is a fast, robust bring-up
path, not the speculative business-intelligence pipeline in
processing/document_processing_engine.py.
"""
# ═══════════════════════════════════════════════════════════════════════
# MODULE: [OPS:ING-001] — shared ingestion helpers (department mapping,
#          chunking, extraction, document-dict construction, backups)
#
# The point of this module existing separately from build_local_index.py
# / build_drive_index.py: both scripts need to chunk/classify/extract
# documents IDENTICALLY, or a document ingested locally in dev would end
# up formatted differently than the same document ingested from Drive in
# prod — every function below is tagged individually (OPS:ING-001a..e).
# ═══════════════════════════════════════════════════════════════════════
from __future__ import annotations

import io
import re
import shutil
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List

# [OPS:ING-001a] FOLDER_DEPARTMENT_MAP / department_for_folder() —
# the folder-name -> department classification table. USED BY:
# [OPS:DRIVE-003] find_finance_folder() callers (build_drive_index.py's
# discover_documents()) and build_local_index.py's discover_documents(),
# via department_for_folder() — a simple substring "in" match, so a
# folder like "Md. Mizanur Rahman (PTIL)" matches the "Md. Mizanur
# Rahman (PTIL)" key even with Drive's exact display name.
# owner folder name (as it appears in Drive / local seed_docs) -> department,
# matching discovery/scanner.py's folder_department_map.
FOLDER_DEPARTMENT_MAP = {
    "Riaz Uddin Sarker": "Admin",
    "Md. Mizanur Rahman (PTIL)": "Finance",
    "Khorshed Alam Babu": "Maintenance",
    "Md. Mozammel Haque": "HR",
    "Zahedul Islam Nizam": "Accounting",
    "Md. Alamin": "Commercial",
    "Monir Ahmed": "Executive",
}


def department_for_folder(folder_name: str) -> str:
    for owner, dept in FOLDER_DEPARTMENT_MAP.items():
        if owner in folder_name:
            return dept
    return "Unknown"


# ─────────────────────────────────────────────────────────────────────────
# [OPS:ING-001b] chunk_text() / extract_text_from_*() / extract_text_from_file()
#                 — paragraph-aware chunking + per-format text extraction
#
# WHAT: chunk_text() packs whole paragraphs into ~900-char windows
#       (never splitting mid-paragraph) rather than a fixed-length
#       sliding window — keeps each chunk semantically coherent, which
#       matters for retrieval quality (a chunk cut mid-sentence embeds
#       worse). extract_text_from_file() dispatches by extension to the
#       format-specific extractor (pypdf for PDF, python-docx for DOCX
#       — including table cells, pandas for XLSX). extract_text_from_
#       xlsx_bytes() flattens sheets to prose; contrast with
#       [OPS:ING-001c] extract_table_chunks_from_xlsx_bytes() which
#       preserves structure instead.
# CALLED BY: make_chunk_documents() [OPS:ING-001d] (chunk_text);
#       build_local_index.py / build_drive_index.py's discover_
#       documents() (extract_text_from_file, directly).
# ─────────────────────────────────────────────────────────────────────────
def chunk_text(text: str, max_chars: int = 900) -> List[str]:
    """Paragraph-aware chunking: pack paragraphs into ~max_chars windows."""
    paragraphs = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    chunks: List[str] = []
    current = ""
    for para in paragraphs:
        if current and len(current) + len(para) + 2 > max_chars:
            chunks.append(current.strip())
            current = para
        else:
            current = f"{current}\n\n{para}" if current else para
    if current:
        chunks.append(current.strip())
    return chunks or [text.strip()]


def extract_text_from_pdf_bytes(raw: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(raw))
    return "\n\n".join((page.extract_text() or "") for page in reader.pages)


def extract_text_from_docx_bytes(raw: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(raw))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n\n".join(parts)


def extract_text_from_xlsx_bytes(raw: bytes, rows_per_group: int = 12) -> str:
    """
    Flatten every sheet into readable text (used for non-table-aware callers).
    Prefer extract_table_chunks_from_xlsx_bytes() when structured row/column
    data needs to survive for downstream table/chart rendering.
    """
    import pandas as pd

    sheets = pd.read_excel(io.BytesIO(raw), sheet_name=None, dtype=str)
    parts: List[str] = []

    for sheet_name, df in sheets.items():
        df = df.dropna(how="all")
        if df.empty:
            continue
        columns = list(df.columns)
        parts.append(f"Sheet '{sheet_name}': columns = {', '.join(str(c) for c in columns)}. Total rows: {len(df)}.")

        records = df.fillna("").to_dict(orient="records")
        for i in range(0, len(records), rows_per_group):
            group = records[i : i + rows_per_group]
            lines = []
            for row in group:
                row_str = ", ".join(f"{k}={v}" for k, v in row.items() if str(v).strip())
                if row_str:
                    lines.append(row_str)
            if lines:
                parts.append(f"Sheet '{sheet_name}' rows {i + 1}-{i + len(group)}:\n" + "\n".join(lines))

    return "\n\n".join(parts)


# ─────────────────────────────────────────────────────────────────────────
# [OPS:ING-001c] extract_table_chunks_from_xlsx_bytes() — the
#                 structure-preserving Excel extractor
#
# WHY THIS EXISTS SEPARATELY FROM extract_text_from_xlsx_bytes(): that
#       function flattens everything to a text blob for embedding/RAG
#       retrieval, losing the actual row/column boundaries. This
#       function returns each row-group as {"text", "table": {"sheet",
#       "columns", "rows"}} instead — the "table" field is what lets the
#       frontend later render a real HTML table or Excel-style chart
#       from a cited chunk, not just quote it as prose.
# CALLED BY: build_local_index.py / build_drive_index.py's discover_
#       documents(), for any .xlsx/.xlsm file — feeds directly into
#       [OPS:ING-001d] make_table_chunk_documents().
# ─────────────────────────────────────────────────────────────────────────
def extract_table_chunks_from_xlsx_bytes(raw: bytes, rows_per_group: int = 12) -> List[Dict[str, Any]]:
    """
    Same row-grouping as extract_text_from_xlsx_bytes(), but returns each
    group as a {"text": ..., "table": {"sheet", "columns", "rows"}} chunk
    instead of a flattened blob — so the actual column/row data survives
    ingestion and can be rendered as a real table or chart later, not just
    read back as prose.
    """
    import pandas as pd

    sheets = pd.read_excel(io.BytesIO(raw), sheet_name=None, dtype=str)
    chunks: List[Dict[str, Any]] = []

    for sheet_name, df in sheets.items():
        df = df.dropna(how="all")
        if df.empty:
            continue
        columns = [str(c) for c in df.columns]
        records = df.fillna("").to_dict(orient="records")

        for i in range(0, len(records), rows_per_group):
            group = records[i : i + rows_per_group]
            lines = [
                f"Sheet '{sheet_name}': columns = {', '.join(columns)}. Rows {i + 1}-{i + len(group)} of {len(records)}."
            ]
            row_values: List[List[str]] = []
            for row in group:
                row_str = ", ".join(f"{k}={v}" for k, v in row.items() if str(v).strip())
                if row_str:
                    lines.append(row_str)
                row_values.append([str(row.get(c, "")) for c in columns])

            chunks.append(
                {
                    "text": "\n".join(lines),
                    "table": {"sheet": sheet_name, "columns": columns, "rows": row_values},
                }
            )

    return chunks


def extract_text_from_file(file_bytes: bytes, file_name: str) -> str:
    """Dispatch by file extension. Returns '' for unsupported types."""
    lower = file_name.lower()
    if lower.endswith(".pdf"):
        return extract_text_from_pdf_bytes(file_bytes)
    if lower.endswith(".docx"):
        return extract_text_from_docx_bytes(file_bytes)
    if lower.endswith((".xlsx", ".xlsm")):
        return extract_text_from_xlsx_bytes(file_bytes)
    if lower.endswith((".txt", ".md")):
        return file_bytes.decode("utf-8", errors="ignore")
    return ""


# ─────────────────────────────────────────────────────────────────────────
# [OPS:ING-001d] make_chunk_documents() / make_table_chunk_documents() —
#                 assembles the final {id, text, metadata} document dicts
#
# WHAT: the text field is prefixed with the document title
#       ("{title}\n\n{chunk}") so the embedding captures the title's
#       context even for a chunk deep in the middle of a long document.
#       metadata carries everything downstream code needs: department
#       (for the salient department badge [OPS:FE-CHAT]), source
#       ("local_seed" vs "google_drive", for the sources display label
#       [OPS:CHAT-012]), chunk_index (used as pgvector's `ordinal`
#       upsert key [OPS:PVEC-002]/[OPS:PVEC-003]).
# CALLED BY: build_local_index.py / build_drive_index.py's discover_
#       documents() — the direct output feeds straight into
#       [OPS:IDX-004] PersistedInMemorySearch.ingest_documents().
# ─────────────────────────────────────────────────────────────────────────
def make_chunk_documents(
    text: str,
    *,
    id_prefix: str,
    file_name: str,
    title: str,
    department: str,
    owner_folder: str,
    source: str,
    extra_metadata: Dict[str, Any] | None = None,
) -> List[Dict[str, Any]]:
    docs: List[Dict[str, Any]] = []
    for i, chunk in enumerate(chunk_text(text)):
        metadata = {
            "source": source,
            "file_name": file_name,
            "title": title,
            "department": department,
            "owner_folder": owner_folder,
            "chunk_index": i,
        }
        if extra_metadata:
            metadata.update(extra_metadata)
        docs.append(
            {
                "id": f"{id_prefix}_{i}",
                "text": f"{title}\n\n{chunk}",
                "metadata": metadata,
            }
        )
    return docs


def make_table_chunk_documents(
    table_chunks: List[Dict[str, Any]],
    *,
    id_prefix: str,
    file_name: str,
    title: str,
    department: str,
    owner_folder: str,
    source: str,
    extra_metadata: Dict[str, Any] | None = None,
) -> List[Dict[str, Any]]:
    """
    Like make_chunk_documents(), but for pre-chunked Excel row-groups (from
    extract_table_chunks_from_xlsx_bytes) where the row/column boundaries are
    already the correct chunk boundaries — no generic paragraph re-chunking.
    """
    docs: List[Dict[str, Any]] = []
    for i, tc in enumerate(table_chunks):
        metadata = {
            "source": source,
            "file_name": file_name,
            "title": title,
            "department": department,
            "owner_folder": owner_folder,
            "chunk_index": i,
            "table": tc["table"],
        }
        if extra_metadata:
            metadata.update(extra_metadata)
        docs.append(
            {
                "id": f"{id_prefix}_{i}",
                "text": f"{title}\n\n{tc['text']}",
                "metadata": metadata,
            }
        )
    return docs


# ─────────────────────────────────────────────────────────────────────────
# [OPS:ING-001e] backup_index_if_exists() — pre-rebuild safety net
#
# WHAT: copies enhanced_index.json to backups/<stem>.<UTC timestamp>.json
#       before ANY operation that's about to unlink() the live index —
#       so a failed/bad reindex can be recovered from instead of losing
#       the corpus outright. Prunes to the most recent `keep` (default
#       5) backups for that specific index file after each call.
# CALLED BY: build_local_index.py main() [--reset flag], build_drive_
#       index.py main() [--reset flag], discovery.py's _run_reindex()
#       [OPS:ADMIN-001a] — every code path that deletes+rebuilds the
#       index calls this immediately before the unlink().
# ─────────────────────────────────────────────────────────────────────────
def backup_index_if_exists(index_path: Path, keep: int = 5) -> Path | None:
    """
    Versioned backup for enhanced_index.json, per the spec's "daily
    versioned backups" requirement — called right before any rebuild wipes
    the file (build_local_index.py --reset, build_drive_index.py --reset,
    and the admin-triggered /api/rag/admin/reindex endpoint all call this).

    Backups live in a sibling `backups/` folder, timestamped, with only the
    most recent `keep` retained. Returns the backup path, or None if there
    was nothing to back up.
    """
    index_path = Path(index_path)
    if not index_path.exists():
        return None

    backups_dir = index_path.parent / "backups"
    backups_dir.mkdir(parents=True, exist_ok=True)

    stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    backup_path = backups_dir / f"{index_path.stem}.{stamp}{index_path.suffix}"
    shutil.copy2(index_path, backup_path)

    # Prune to the most recent `keep` backups for this index file.
    existing = sorted(backups_dir.glob(f"{index_path.stem}.*{index_path.suffix}"), key=lambda p: p.stat().st_mtime, reverse=True)
    for stale in existing[keep:]:
        stale.unlink(missing_ok=True)

    return backup_path
